import io

import pytest
from docx import Document
from openpyxl import Workbook

from whisp.chat.documents import MAX_TEXT_CHARS, UnsupportedDocument, extract


def make_pdf(text: str) -> bytes:
    """Build a minimal one-page PDF with a real text layer."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def test_pdf_text_is_extracted() -> None:
    result = extract("invoice.pdf", "application/octet-stream", make_pdf("Total due 1250 USD"))

    assert "Total due 1250 USD" in result.text


def test_docx_paragraphs_and_tables_are_extracted() -> None:
    document = Document()
    document.add_paragraph("Contract renewal terms")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Price"
    table.rows[0].cells[1].text = "500"
    buffer = io.BytesIO()
    document.save(buffer)

    result = extract("terms.docx", "application/octet-stream", buffer.getvalue())

    assert "Contract renewal terms" in result.text
    assert "Price | 500" in result.text


def test_xlsx_rows_are_extracted_per_sheet() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Cost"])
    sheet.append(["Laptop", 1200])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = extract("budget.xlsx", "application/octet-stream", buffer.getvalue())

    assert "## Sheet: Budget" in result.text
    assert "Laptop | 1200" in result.text


def test_csv_and_html_are_converted_to_text() -> None:
    assert "a | b" in extract("data.csv", "text/csv", b"a,b\n1,2").text
    html = extract("page.html", "text/html", b"<p>Hello <b>there</b></p><script>x()</script>")
    assert "Hello" in html.text and "x()" not in html.text


def test_images_are_passed_through_for_vision() -> None:
    result = extract("photo.jpg", "image/jpeg", b"\xff\xd8fake")

    assert result.images == (("image/jpeg", b"\xff\xd8fake"),)


def test_long_text_is_truncated() -> None:
    result = extract("notes.txt", "text/plain", b"x" * (MAX_TEXT_CHARS + 500))

    assert result.text.endswith("[... truncated ...]")


def test_unsupported_and_oversized_files_are_refused() -> None:
    with pytest.raises(UnsupportedDocument):
        extract("archive.zip", "application/zip", b"PK")
    with pytest.raises(UnsupportedDocument):
        extract("big.png", "image/png", b"0" * (5 * 1024 * 1024 + 1))
