import csv
import io
from pathlib import PurePosixPath

from whisp.chat.models import ToolResult
from whisp.core.errors import SafeWhispError
from whisp.text import html_to_text

MAX_TEXT_CHARS = 20_000
MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_SHEETS = 10
MAX_ROWS_PER_SHEET = 200

IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/plain",
    ".htm": "text/html",
    ".html": "text/html",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}


class UnsupportedDocument(SafeWhispError):
    """The document type cannot be read, or its content could not be extracted."""


def document_type(filename: str, mime_type: str) -> str:
    # Mail clients often label attachments application/octet-stream, so the extension
    # is the more reliable signal.
    by_extension = _EXTENSIONS.get(PurePosixPath(filename.lower()).suffix)
    return by_extension or mime_type.split(";")[0].strip().lower()


def extract(filename: str, mime_type: str, data: bytes) -> ToolResult:
    """Turn a document into text, or into an image for a vision model."""
    kind = document_type(filename, mime_type)
    if kind in IMAGE_TYPES:
        if len(data) > MAX_IMAGE_BYTES:
            raise UnsupportedDocument("Image is too large to analyze (limit 5 MB)")
        return ToolResult(text=f"Image {filename} is attached below.", images=((kind, data),))
    if kind == "application/pdf":
        text = _pdf(data)
    elif kind.endswith("wordprocessingml.document"):
        text = _docx(data)
    elif kind.endswith("spreadsheetml.sheet"):
        text = _xlsx(data)
    elif kind == "text/csv":
        text = _csv(data)
    elif kind == "text/html":
        text = html_to_text(_decode(data))
    elif kind.startswith("text/"):
        text = _decode(data)
    else:
        raise UnsupportedDocument(f"Cannot read files of type {kind}")
    return ToolResult(text=_limit(text))


def _decode(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _limit(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS] + "\n[... truncated ...]"
    return text


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnsupportedDocument("PDF is password-protected")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError:
        raise UnsupportedDocument("PDF could not be read") from None
    text = "\n\n".join(page.strip() for page in pages if page.strip())
    if not text:
        # Scanned PDFs are images of text; say so rather than returning nothing.
        return "(This PDF has no text layer; it is probably a scanned image.)"
    return text


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception:
        raise UnsupportedDocument("Word document could not be read") from None
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise UnsupportedDocument("Excel workbook could not be read") from None
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            parts.append(f"## Sheet: {sheet.title}")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS_PER_SHEET:
                    parts.append(f"[... rows after {MAX_ROWS_PER_SHEET} omitted ...]")
                    break
                if any(value is not None for value in row):
                    parts.append(" | ".join("" if value is None else str(value) for value in row))
    finally:
        workbook.close()
    return "\n".join(parts)


def _csv(data: bytes) -> str:
    rows = csv.reader(io.StringIO(_decode(data)))
    return "\n".join(" | ".join(row) for row in rows)
